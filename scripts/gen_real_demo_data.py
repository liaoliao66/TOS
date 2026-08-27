# -*- coding: utf-8 -*-
"""从《每日每班时间作业统计表》生成原型演示数据（支持多文件合并）。

在线演示推荐流程：
  1. 将本周 Excel 放入 文件/ 目录
  2. git add + git push
  3. GitHub Actions 自动合并全部 xlsx → 更新 real-work-stat-demo.js → Pages 刷新

本地：
  python scripts/gen_real_demo_data.py           # 默认合并 文件/ 下全部 xlsx
  python scripts/gen_real_demo_data.py --latest  # 仅使用最新一个 xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "文件"
OUT_JSON = ROOT / "prototype" / "versions" / "v1.0.0" / "assets" / "real-work-stat-demo.json"
OUT_JS = ROOT / "prototype" / "versions" / "v1.0.0" / "assets" / "real-work-stat-demo.js"

L1_MAP = {
    "工铵吨包": "吨包袋",
    "氧化钙": "散货",
    "氮磷肥": "散货",
    "硫矿": "散货",
    "磷矿": "散货",
    "二氢钾": "散货",
    "吨包": "吨包袋",
    "硫磺": "散货",
    "磷酸二氢氨": "散货",
    "脱硫石膏": "散货",
}


def excel_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and v > 30000:
        return date(1899, 12, 30) + timedelta(days=int(v))
    return None


def norm_cargo(v):
    if not isinstance(v, str):
        return None
    s = v.strip()
    if not s:
        return None
    if s.replace("，", ",") == "二氢钾,氮磷肥":
        return "二氢钾"
    return s


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return None


def is_period(v):
    if not isinstance(v, str):
        return False
    return bool(re.match(r"^\d{1,2}:\d{2}\s*-\s*\d{1,2}:\d{2}$", v.strip()))


def parse_sheet(ws):
    berth_starts = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(2, c).value
        if v and "泊位" in str(v):
            berth_starts.append((c, str(v).strip()))

    blocks = []
    for i, (start, name) in enumerate(berth_starts):
        end = berth_starts[i + 1][0] - 1 if i + 1 < len(berth_starts) else ws.max_column
        machines = []
        c = start
        while c <= end:
            if ws.cell(3, c).value == "机械名称":
                fields = {"机械名称": c}
                cc = c + 1
                while cc <= end:
                    h = ws.cell(3, cc).value
                    if h == "机械名称":
                        break
                    if h:
                        fields[str(h).strip()] = cc
                    cc += 1
                machines.append(fields)
                c = cc
            else:
                c += 1
        blocks.append({"berth": name, "machines": machines})

    records = []
    last_date = None
    last_shift = None
    cargo_ff = {}

    for r in range(4, ws.max_row + 1):
        d = excel_date(ws.cell(r, 1).value)
        shift_raw = ws.cell(r, 2).value
        period = ws.cell(r, 3).value

        if d:
            if last_date and d != last_date:
                cargo_ff = {}
            last_date = d

        if isinstance(shift_raw, str):
            s = shift_raw.strip()
            if s in ("白班", "夜班"):
                last_shift = s

        if not last_date or not last_shift:
            continue
        if not is_period(period):
            continue

        for bi, b in enumerate(blocks):
            for mi, m in enumerate(b["machines"]):
                key = (bi, mi)
                driver = ws.cell(r, m["司机姓名"]).value if "司机姓名" in m else None
                if isinstance(driver, str):
                    driver = driver.strip() or None
                elif driver is not None:
                    driver = str(driver).strip() or None

                cargo_raw = ws.cell(r, m["货种"]).value if "货种" in m else None
                cargo = norm_cargo(cargo_raw)
                if cargo:
                    cargo_ff[key] = cargo
                else:
                    cargo = cargo_ff.get(key)

                qty = to_num(ws.cell(r, m["作业量"]).value) if "作业量" in m else None
                machine = ws.cell(r, m["机械名称"]).value if "机械名称" in m else None
                if isinstance(machine, str):
                    machine = machine.strip() or None
                remark = ws.cell(r, m["备注（原因）"]).value if "备注（原因）" in m else None
                if isinstance(remark, str):
                    remark = remark.strip() or None

                if qty is None or qty == 0:
                    continue
                if not driver or not cargo:
                    continue

                records.append(
                    {
                        "date": last_date.isoformat(),
                        "shift": last_shift,
                        "period": str(period).strip().replace(" ", ""),
                        "berth": b["berth"],
                        "machine": machine,
                        "driver": driver,
                        "cargo": cargo,
                        "qty": round(qty, 2),
                        "remark": remark,
                        "abnormal": bool(remark),
                    }
                )
    return records


def record_key(r: dict) -> tuple:
    return (
        r["date"],
        r["shift"],
        r["period"],
        r["berth"],
        r["machine"],
        r["driver"],
        r["cargo"],
        r["qty"],
        r.get("remark"),
    )


def load_records(paths: list[Path]) -> list[dict]:
    from openpyxl import load_workbook

    records: list[dict] = []
    seen: set[tuple] = set()
    for path in paths:
        wb = load_workbook(path, data_only=True)
        for name in wb.sheetnames:
            for r in parse_sheet(wb[name]):
                key = record_key(r)
                if key in seen:
                    continue
                seen.add(key)
                records.append(r)
    return records


def list_xlsx() -> list[Path]:
    if not DATA_DIR.is_dir():
        sys.exit(f"找不到数据目录：{DATA_DIR}")
    files = sorted(DATA_DIR.glob("*.xlsx"), key=lambda p: p.stat().st_mtime)
    if not files:
        sys.exit(f"在 {DATA_DIR} 下未找到 .xlsx，请将统计表放入该目录后 push")
    return files


def resolve_inputs(mode: str, explicit: Path | None) -> list[Path]:
    if explicit:
        p = explicit if explicit.is_absolute() else ROOT / explicit
        if not p.is_file():
            sys.exit(f"找不到 Excel：{p}")
        return [p]
    files = list_xlsx()
    if mode == "latest":
        return [files[-1]]
    return files


def build_payload(records: list[dict], sources: list[str]) -> dict:
    if not records:
        sys.exit("未解析到任何作业记录，请检查 Excel 表结构")

    dates = sorted(set(r["date"] for r in records))
    drivers = sorted(set(r["driver"] for r in records))
    cargos = sorted(set(r["cargo"] for r in records))

    by_cargo_day = {c: {d: 0.0 for d in dates} for c in cargos}
    by_driver_cargo_day = {c: {dr: {d: 0.0 for d in dates} for dr in drivers} for c in cargos}
    hours_driver_shift_cargo = defaultdict(float)
    hours_driver_shift = defaultdict(float)
    qty_driver_shift_cargo = defaultdict(lambda: {"total": 0.0, "normal": 0.0, "abnormal": 0.0})
    qty_driver_shift = defaultdict(lambda: {"total": 0.0, "normal": 0.0, "abnormal": 0.0})
    seen_hour: set[tuple] = set()
    seen_hour_all: set[tuple] = set()
    by_day_shift = defaultdict(lambda: {"day": 0.0, "night": 0.0})
    by_cargo_day_shift = defaultdict(lambda: defaultdict(lambda: {"day": 0.0, "night": 0.0}))
    ab_counter = Counter()
    normal_qty = 0.0
    ab_qty = 0.0

    for r in records:
        by_cargo_day[r["cargo"]][r["date"]] += r["qty"]
        by_driver_cargo_day[r["cargo"]][r["driver"]][r["date"]] += r["qty"]
        k = (r["driver"], r["shift"], r["cargo"], r["date"], r["period"])
        k2 = (r["driver"], r["shift"], r["date"], r["period"])
        if k not in seen_hour:
            seen_hour.add(k)
            hours_driver_shift_cargo[(r["driver"], r["shift"], r["cargo"])] += 1.0
        if k2 not in seen_hour_all:
            seen_hour_all.add(k2)
            hours_driver_shift[(r["driver"], r["shift"])] += 1.0
        bucket = qty_driver_shift_cargo[(r["driver"], r["shift"], r["cargo"])]
        bucket["total"] += r["qty"]
        if r["abnormal"]:
            bucket["abnormal"] += r["qty"]
            ab_qty += r["qty"]
            ab_counter[r["remark"] or "其他"] += r["qty"]
        else:
            bucket["normal"] += r["qty"]
            normal_qty += r["qty"]
        bucket2 = qty_driver_shift[(r["driver"], r["shift"])]
        bucket2["total"] += r["qty"]
        if r["abnormal"]:
            bucket2["abnormal"] += r["qty"]
        else:
            bucket2["normal"] += r["qty"]
        key = "day" if r["shift"] == "白班" else "night"
        by_day_shift[r["date"]][key] += r["qty"]
        by_cargo_day_shift[r["cargo"]][r["date"]][key] += r["qty"]

    evidence = {}
    for c in cargos:
        masters = {}
        for dr in drivers:
            series = [round(by_driver_cargo_day[c][dr][d], 2) for d in dates]
            if sum(series) > 0:
                masters[dr] = series
        flat = [v for series in masters.values() for v in series if v > 0]
        evidence[c] = {
            "avg": round(sum(flat) / len(flat), 2) if flat else 0,
            "max": round(max(flat), 2) if flat else 0,
            "min": round(min(flat), 2) if flat else 0,
            "masters": masters,
        }

    volume_by_cargo_day = {c: [round(by_cargo_day[c][d], 2) for d in dates] for c in cargos}

    def make_rank(scope_filter):
        rows = []
        if scope_filter is None:
            for (dr, sh), q in sorted(qty_driver_shift.items()):
                h = hours_driver_shift[(dr, sh)]
                rows.append(
                    {
                        "name": dr,
                        "shift": sh,
                        "hours": round(h, 2),
                        "totalQty": round(q["total"], 2),
                        "normalQty": round(q["normal"], 2),
                        "abnormalQty": round(q["abnormal"], 2),
                    }
                )
        elif scope_filter[0] == "l2":
            cargo = scope_filter[1]
            for (dr, sh, c), q in qty_driver_shift_cargo.items():
                if c != cargo:
                    continue
                h = hours_driver_shift_cargo[(dr, sh, c)]
                rows.append(
                    {
                        "name": dr,
                        "shift": sh,
                        "hours": round(h, 2),
                        "totalQty": round(q["total"], 2),
                        "normalQty": round(q["normal"], 2),
                        "abnormalQty": round(q["abnormal"], 2),
                    }
                )
        elif scope_filter[0] == "l1":
            l1 = scope_filter[1]
            kids = [c for c, p in L1_MAP.items() if p == l1]
            agg = defaultdict(lambda: {"total": 0.0, "normal": 0.0, "abnormal": 0.0})
            hours = defaultdict(float)
            seen = set()
            for r in records:
                if r["cargo"] not in kids:
                    continue
                agg[(r["driver"], r["shift"])]["total"] += r["qty"]
                if r["abnormal"]:
                    agg[(r["driver"], r["shift"])]["abnormal"] += r["qty"]
                else:
                    agg[(r["driver"], r["shift"])]["normal"] += r["qty"]
                kk = (r["driver"], r["shift"], r["date"], r["period"])
                if kk not in seen:
                    seen.add(kk)
                    hours[(r["driver"], r["shift"])] += 1.0
            for (dr, sh), q in sorted(agg.items()):
                rows.append(
                    {
                        "name": dr,
                        "shift": sh,
                        "hours": round(hours[(dr, sh)], 2),
                        "totalQty": round(q["total"], 2),
                        "normalQty": round(q["normal"], 2),
                        "abnormalQty": round(q["abnormal"], 2),
                    }
                )
        return rows

    def kpi_from_rows(rows):
        if not rows:
            return {"total": "0.00", "avg": "0.00", "max": "0.00", "min": "0.00"}
        total = sum(r["totalQty"] for r in rows)
        avgs = [(r["normalQty"] / r["hours"] if r["hours"] else 0) for r in rows]
        return {
            "total": f"{total:,.2f}",
            "avg": f"{sum(avgs) / len(avgs):.2f}",
            "max": f"{max(avgs):.2f}",
            "min": f"{min(avgs):.2f}",
        }

    rank_all = make_rank(None)
    rank_l1 = {l1: make_rank(("l1", l1)) for l1 in ["吨包袋", "散货"]}
    rank_l2 = {c: make_rank(("l2", c)) for c in cargos}
    kpi = {"全部": kpi_from_rows(rank_all)}
    for l1, rows in rank_l1.items():
        kpi[l1] = kpi_from_rows(rows)
    for c, rows in rank_l2.items():
        kpi[c] = kpi_from_rows(rows)

    volume_day_shift = [
        {"date": d, "day": round(by_day_shift[d]["day"], 2), "night": round(by_day_shift[d]["night"], 2)}
        for d in dates
    ]
    volume_cargo_day_shift = {
        c: [
            {
                "date": d,
                "day": round(by_cargo_day_shift[c][d]["day"], 2),
                "night": round(by_cargo_day_shift[c][d]["night"], 2),
            }
            for d in dates
        ]
        for c in cargos
    }

    cargo_totals = {c: round(sum(by_cargo_day[c].values()), 2) for c in cargos}
    cargo_shift = {}
    for c in cargos:
        day = sum(by_cargo_day_shift[c][d]["day"] for d in dates)
        night = sum(by_cargo_day_shift[c][d]["night"] for d in dates)
        cargo_shift[c] = {"day": round(day, 2), "night": round(night, 2), "total": round(day + night, 2)}

    cargo_rank = []
    for c in sorted(cargos, key=lambda x: -cargo_shift[x]["total"]):
        cargo_rank.append(
            {
                "name": c,
                "l1": L1_MAP.get(c, "散货"),
                "day": cargo_shift[c]["day"],
                "night": cargo_shift[c]["night"],
            }
        )

    volume_shift_by_cargo = {
        c: {
            "day": [round(by_cargo_day_shift[c][d]["day"], 2) for d in dates],
            "night": [round(by_cargo_day_shift[c][d]["night"], 2) for d in dates],
        }
        for c in cargos
    }

    evidence_days_md = [d[5:] for d in dates]

    def hydrate_rank(rows):
        out_rows = []
        for r in rows:
            h = r["hours"] or 0
            out_rows.append(
                {
                    "name": r["name"],
                    "shift": r["shift"],
                    "hours": r["hours"],
                    "totalQty": r["totalQty"],
                    "totalAvg": (r["totalQty"] / h) if h else 0,
                    "normalQty": r["normalQty"],
                    "normalAvg": (r["normalQty"] / h) if h else 0,
                    "abnormalQty": r["abnormalQty"],
                    "abnormalAvg": (r["abnormalQty"] / h) if h else 0,
                }
            )
        return out_rows

    source_label = sources[-1] if len(sources) == 1 else f"{len(sources)} 个文件（最新：{sources[-1]}）"

    return {
        "meta": {
            "source": source_label,
            "sources": sources,
            "date_from": dates[0],
            "date_to": dates[-1],
            "n_records": len(records),
            "drivers": drivers,
            "cargos": cargos,
            "l1_map": L1_MAP,
            "total_qty": round(sum(r["qty"] for r in records), 2),
            "normal_qty": round(normal_qty, 2),
            "abnormal_qty": round(ab_qty, 2),
            "normal_ratio": round(normal_qty / (normal_qty + ab_qty), 4) if (normal_qty + ab_qty) else 0,
        },
        "EVIDENCE_DAYS": evidence_days_md,
        "EVIDENCE_DATES": dates,
        "EVIDENCE_DATA": evidence,
        "VOLUME_DAYS": evidence_days_md,
        "VOLUME_BY_CARGO_DAY": volume_by_cargo_day,
        "VOLUME_SHIFT_BY_CARGO": volume_shift_by_cargo,
        "VOLUME_DAY_SHIFT": volume_day_shift,
        "VOLUME_CARGO_DAY_SHIFT": volume_cargo_day_shift,
        "VOLUME_DAY_ROWS_ALL": list(reversed(volume_day_shift)),
        "VOLUME_DAY_ROWS_BY_CARGO": {c: list(reversed(volume_cargo_day_shift[c])) for c in cargos},
        "CARGO_RANK": cargo_rank,
        "CARGO_TOTALS": cargo_totals,
        "CARGO_SHIFT": cargo_shift,
        "DRIVER_OPTIONS": drivers,
        "RANK_ROWS_ALL": hydrate_rank(rank_all),
        "RANK_BY_L1": {k: hydrate_rank(v) for k, v in rank_l1.items()},
        "RANK_BY_L2": {k: hydrate_rank(v) for k, v in rank_l2.items()},
        "KPI_BY_SCOPE": kpi,
        "ABNORMAL_TYPE_TONS": [{"name": k, "tons": round(v, 2)} for k, v in ab_counter.most_common(10)],
        "L2_COLORS": {
            "氧化钙": "#0f766e",
            "氮磷肥": "#2563eb",
            "硫矿": "#7c3aed",
            "磷矿": "#0891b2",
            "二氢钾": "#ca8a04",
            "工铵吨包": "#db2777",
        },
        "MASTER_COLORS": {
            "涂峰": "#0f766e",
            "陈迪政": "#2563eb",
            "刘志显": "#7c3aed",
            "雷雨": "#dc2626",
            "柯力": "#d97706",
            "刘念": "#059669",
            "胡涛": "#4f46e5",
            "聂星宇": "#db2777",
            "李军民": "#0ea5e9",
            "杨湘岳": "#65a30d",
        },
    }


def write_outputs(payload: dict) -> None:
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    js = (
        "/* auto-generated from Excel — do not edit by hand */\n"
        "window.RealWorkStatDemo = "
        + json.dumps(payload, ensure_ascii=False)
        + ";\n"
    )
    OUT_JS.write_text(js, encoding="utf-8")


def generate(paths: list[Path]) -> dict:
    records = load_records(paths)
    sources = [p.name for p in paths]
    payload = build_payload(records, sources)
    write_outputs(payload)
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(description="Excel → 原型演示数据")
    parser.add_argument("--file", type=Path, help="指定单个 Excel")
    parser.add_argument(
        "--latest",
        action="store_true",
        help="仅使用 文件/ 下最新 xlsx（默认合并全部 xlsx）",
    )
    args = parser.parse_args()

    mode = "latest" if args.latest else "merge"
    paths = resolve_inputs(mode, args.file)
    payload = generate(paths)
    meta = payload["meta"]

    print("—" * 48)
    print(f"输入：{len(paths)} 个文件")
    for name in meta.get("sources", []):
        print(f"  · {name}")
    print(
        f"记录 {meta['n_records']} 条 · {meta['date_from']} ~ {meta['date_to']} · "
        f"总产量 {meta['total_qty']} 吨"
    )
    print(f"已写入：{OUT_JS.relative_to(ROOT)}")
    print("—" * 48)


if __name__ == "__main__":
    main()
